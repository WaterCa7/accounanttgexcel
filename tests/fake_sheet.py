"""Таблица в памяти, которая отвечает как Google Sheets API. Нужна только для тестов."""

from __future__ import annotations

import re
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from table import parse_range

EPOCH = date(1899, 12, 30)
RANGE_RE = re.compile(r"^'((?:[^']|'')+)'(?:!(.+))?$")
SUM_RE = re.compile(r"^=SUM\(([A-Z]+\d+:[A-Z]+\d+)\)$", re.I)
ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class FakeBook:
    def __init__(self, title="Тестовая таблица"):
        self.title = title
        self.url = "https://docs.google.com/spreadsheets/d/test-id"
        self.sheets: list[dict] = []
        self.requests: list[tuple[str, object]] = []  # что бот отправлял на запись

    def add_sheet(self, title: str, rows: int = 100, cols: int = 15) -> dict:
        sheet = {"id": len(self.sheets) + 10, "title": title, "rows": rows, "cols": cols, "cells": {}}
        self.sheets.append(sheet)
        return sheet

    @classmethod
    def from_xlsx(cls, path) -> "FakeBook":
        import openpyxl

        book = cls(Path(path).stem)
        for source in openpyxl.load_workbook(path).worksheets:
            sheet = book.add_sheet(source.title, max(source.max_row, 100), max(source.max_column, 15))
            for row in source.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        value = cell.value.date() if isinstance(cell.value, datetime) else cell.value
                        sheet["cells"][(cell.row, cell.column)] = value
        return book

    def cell(self, title: str, address: str):
        row, col, _, _ = parse_range(address)
        return self._sheet(title)["cells"].get((row, col), "")

    # --- то, что вызывает table.py ---

    def fetch_sheet_metadata(self, params=None) -> dict:
        return {
            "properties": {"title": self.title, "locale": "ru_RU", "timeZone": "Europe/Moscow"},
            "sheets": [{"properties": {
                "sheetId": s["id"], "title": s["title"],
                "gridProperties": {"rowCount": s["rows"], "columnCount": s["cols"]},
            }} for s in self.sheets],
        }

    def values_batch_get(self, ranges, params=None) -> dict:
        render = (params or {}).get("valueRenderOption", "FORMATTED_VALUE")
        return {"valueRanges": [self._values(r, render) for r in ranges]}

    def values_batch_update(self, body) -> dict:
        assert body["valueInputOption"] == "USER_ENTERED"
        self.requests.append(("values", body["data"]))
        for item in body["data"]:
            sheet, (row1, col1, _, _) = self._locate(item["range"])
            for dr, values in enumerate(item["values"]):
                for dc, value in enumerate(values):
                    key = (row1 + dr, col1 + dc)
                    if value == "":
                        sheet["cells"].pop(key, None)
                    else:
                        sheet["cells"][key] = self._entered(value)
        return {}

    def batch_update(self, body) -> dict:
        self.requests.append(("batch", body))
        for request in body["requests"]:
            (kind, payload), = request.items()
            rng = payload["range"]
            assert rng["dimension"] == "ROWS"
            sheet = next(s for s in self.sheets if s["id"] == rng["sheetId"])
            start, count = rng["startIndex"] + 1, rng["endIndex"] - rng["startIndex"]
            moved = {}
            for (row, col), value in sheet["cells"].items():
                if kind == "insertDimension":
                    moved[(row + count if row >= start else row, col)] = value
                elif kind == "deleteDimension":
                    if row >= start + count:
                        moved[(row - count, col)] = value
                    elif row < start:
                        moved[(row, col)] = value
                else:
                    raise AssertionError(kind)
            sheet["cells"] = moved
            sheet["rows"] += count if kind == "insertDimension" else -count
        return {}

    # --- внутреннее ---

    def _sheet(self, title: str) -> dict:
        return next(s for s in self.sheets if s["title"] == title)

    def _locate(self, address: str):
        match = RANGE_RE.match(address)
        assert match, address
        sheet = self._sheet(match.group(1).replace("''", "'"))
        if match.group(2):
            return sheet, parse_range(match.group(2))
        last_row = max((r for r, _ in sheet["cells"]), default=0)
        last_col = max((c for _, c in sheet["cells"]), default=0)
        return sheet, (1, 1, last_row, last_col)

    def _values(self, address: str, render: str) -> dict:
        sheet, (row1, col1, row2, col2) = self._locate(address)
        grid = []
        for row in range(row1, row2 + 1):
            line = [self._render(sheet, sheet["cells"].get((row, col), ""), render) for col in range(col1, col2 + 1)]
            while line and line[-1] == "":
                line.pop()
            grid.append(line)
        while grid and not grid[-1]:
            grid.pop()
        return {"range": address, "values": grid} if grid else {"range": address}

    def _render(self, sheet, value, render: str):
        if render == "FORMULA":
            return (value - EPOCH).days if isinstance(value, date) else value
        if isinstance(value, date):
            return value.strftime("%d-%b-%y").lstrip("0")
        if isinstance(value, str) and value.startswith("="):
            match = SUM_RE.match(value)
            if not match:
                return "#N/A"
            row1, col1, row2, col2 = parse_range(match.group(1))
            numbers = [
                v for (r, c), v in sheet["cells"].items()
                if row1 <= r <= row2 and col1 <= c <= col2 and isinstance(v, (int, float)) and not isinstance(v, bool)
            ]
            return f"{sum(numbers):,.2f}"
        if isinstance(value, float):
            return f"{value:,.2f}"
        return str(value) if value != "" else ""

    @staticmethod
    def _entered(value):
        """Как таблица понимает ввод с клавиатуры (USER_ENTERED)."""
        if isinstance(value, str):
            if ISO_RE.match(value):
                return date.fromisoformat(value)
            try:
                return float(value) if "." in value else int(value)
            except ValueError:
                return value
        return value
